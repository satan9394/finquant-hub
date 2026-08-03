"""Versioned research-report export helpers.

The database remains the single source of truth: HTML, Markdown, PDF and
DOCX are rendered on demand from one saved research-report version.
"""
from __future__ import annotations

from html import escape
from io import BytesIO
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer


SECTIONS = (
    ("执行摘要", "executive_summary"),
    ("商业模式", "business_model_analysis"),
    ("财务分析", "financial_analysis"),
    ("行业分析", "industry_analysis"),
    ("风险提示", "risk_analysis"),
    ("技术分析说明", "technical_analysis"),
    ("综合结论", "verdict"),
)
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_URL_RE = re.compile(r"https?://[^\s\]\[\"'<>]+")


def report_to_markdown(report: dict[str, Any]) -> str:
    """Create a portable Markdown representation of one persisted version."""
    data = report.get("report_data") or {}
    company = data.get("company_name") or report.get("company_name") or report.get("ticker")
    ticker = data.get("ticker") or report.get("ticker")
    created_at = report.get("created_at") or ""
    model = f"模型 #{report['model_id']}" if report.get("model_id") else "默认模型"
    lines = [
        f"# {company} 研究报告",
        "",
        f"> 股票代码：{ticker}  ",
        f"> 报告版本：#{report.get('id')}  ",
        f"> 生成时间：{created_at}  ",
        f"> 生成模型：{model}",
    ]
    chapters = data.get("chapters")
    if isinstance(chapters, list):
        for chapter in chapters:
            if not isinstance(chapter, dict) or not chapter.get("content"):
                continue
            lines.extend(["", f"## {chapter.get('title') or chapter.get('key')}", "", str(chapter["content"]).strip()])
    else:
        for title, key in SECTIONS:
            value = data.get(key)
            if value:
                lines.extend(["", f"## {title}", "", str(value).strip()])

    urls = _source_urls(report.get("agent_results") or {})
    if urls:
        lines.extend(["", "## 来源与证据", ""])
        lines.extend(f"- {url}" for url in urls)

    if len(lines) == 6:
        lines.extend(["", "报告尚未生成完整正文。"])
    return "\n".join(lines).strip() + "\n"


def markdown_to_html(markdown: str, title: str) -> str:
    """Render the conservative Markdown subset used by research reports safely."""
    body: list[str] = []
    list_open = False
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if line.startswith("- "):
            if not list_open:
                body.append("<ul>")
                list_open = True
            body.append(f"<li>{_inline_html(line[2:])}</li>")
            continue
        if list_open:
            body.append("</ul>")
            list_open = False
        if not line:
            continue
        if line == "---":
            body.append("<hr>")
        elif line.startswith("### "):
            body.append(f"<h3>{_inline_html(line[4:])}</h3>")
        elif line.startswith("## "):
            body.append(f"<h2>{_inline_html(line[3:])}</h2>")
        elif line.startswith("# "):
            body.append(f"<h1>{_inline_html(line[2:])}</h1>")
        elif line.startswith("> "):
            body.append(f"<p class=\"meta\">{_inline_html(line[2:])}</p>")
        else:
            body.append(f"<p>{_inline_html(line)}</p>")
    if list_open:
        body.append("</ul>")

    safe_title = escape(title)
    return f"""<!doctype html>
<html lang=\"zh-CN\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">
<title>{safe_title}</title><style>
body {{ max-width: 860px; margin: 40px auto; padding: 0 24px; color: #1f2937; font: 16px/1.75 -apple-system, BlinkMacSystemFont, \"Microsoft YaHei\", sans-serif; }}
h1 {{ color: #0f3d75; font-size: 30px; margin-bottom: 24px; }} h2 {{ color: #0f3d75; border-bottom: 1px solid #dbe5f1; padding-bottom: 6px; margin-top: 30px; }}
h3 {{ color: #244b74; }} .meta {{ color: #64748b; margin: 2px 0; }} li {{ margin: 4px 0; }}
</style></head><body>{''.join(body)}</body></html>"""


def markdown_to_docx(markdown: str, title: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].font.color.rgb = RGBColor(15, 61, 117)
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("# "):
            continue
        if line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_docx_runs(paragraph, line[2:])
        elif line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(line[2:])
            run.italic = True
            run.font.color.rgb = RGBColor(100, 116, 139)
        else:
            paragraph = document.add_paragraph()
            _add_docx_runs(paragraph, line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("FinQuant Hub · AI 生成内容不构成投资建议").font.size = Pt(8)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def markdown_to_pdf(markdown: str, title: str) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    base = ParagraphStyle("ResearchBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=10.5, leading=17, spaceAfter=7)
    title_style = ParagraphStyle("ResearchTitle", parent=base, fontSize=22, leading=28, alignment=TA_CENTER, textColor=colors.HexColor("#0f3d75"), spaceAfter=16)
    h2 = ParagraphStyle("ResearchH2", parent=base, fontSize=15, leading=22, textColor=colors.HexColor("#0f3d75"), spaceBefore=14, spaceAfter=8)
    h3 = ParagraphStyle("ResearchH3", parent=base, fontSize=12, leading=19, textColor=colors.HexColor("#244b74"), spaceBefore=10)
    meta = ParagraphStyle("ResearchMeta", parent=base, fontSize=9, leading=13, textColor=colors.HexColor("#64748b"), spaceAfter=1)
    story: list[Flowable] = [Paragraph(escape(title), title_style)]
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("# "):
            continue
        if line.startswith("## "):
            story.append(Paragraph(_inline_html(line[3:]), h2))
        elif line.startswith("### "):
            story.append(Paragraph(_inline_html(line[4:]), h3))
        elif line.startswith("> "):
            story.append(Paragraph(_inline_html(line[2:]), meta))
        elif line.startswith("- "):
            story.append(Paragraph(f"• {_inline_html(line[2:])}", base))
        else:
            story.append(Paragraph(_inline_html(line), base))
    story.append(Spacer(1, 6))
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return stream.getvalue()


def _inline_html(value: str) -> str:
    safe = escape(value)
    return _BOLD_RE.sub(r"<b>\1</b>", safe)


def _add_docx_runs(paragraph: Any, value: str) -> None:
    position = 0
    for match in _BOLD_RE.finditer(value):
        if match.start() > position:
            paragraph.add_run(value[position:match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        position = match.end()
    if position < len(value):
        paragraph.add_run(value[position:])


def _source_urls(value: Any) -> list[str]:
    found: set[str] = set()
    def walk(item: Any) -> None:
        if isinstance(item, dict):
            for child in item.values():
                walk(child)
        elif isinstance(item, list):
            for child in item:
                walk(child)
        elif isinstance(item, str):
            found.update(_URL_RE.findall(item))
    walk(value)
    return sorted(found)


def _pdf_footer(canvas: Any, doc: Any) -> None:
    canvas.saveState()
    canvas.setFont("STSong-Light", 8)
    canvas.setFillColor(colors.HexColor("#64748b"))
    canvas.drawCentredString(A4[0] / 2, 10 * mm, f"FinQuant Hub · 第 {doc.page} 页 · AI 生成内容不构成投资建议")
    canvas.restoreState()
