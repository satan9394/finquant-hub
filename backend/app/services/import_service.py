"""Watchlist text/file parsing and instrument resolution."""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from collections.abc import Iterable
from typing import Any, Protocol

from openpyxl import load_workbook

from .market_data import normalize_code


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_CANDIDATES = 500
_HEADERS = {
    "代码", "股票代码", "证券代码", "名称", "股票名称", "证券名称",
    "code", "symbol", "ticker", "name",
}
_CODE_HEADERS = {"代码", "股票代码", "证券代码", "code", "symbol", "ticker"}
_NAME_HEADERS = {"名称", "股票名称", "证券名称", "name"}
_DIRECT_CODE = re.compile(
    r"(?i)(?:\d{4,6}|(?:sh|sz)\d{6}|hk\d{4,5}|us[A-Z][A-Z0-9.-]{0,14}|"
    r"\d{4,6}\.(?:HK|SS|SZ)|[A-Z][A-Z0-9.-]{0,14})"
)


class ImportFormatError(ValueError):
    """Raised when an uploaded watchlist file cannot be parsed safely."""


class SearchMarket(Protocol):
    async def search(self, keyword: str) -> list[dict[str, Any]]: ...


def _clean_candidates(values: Iterable[object]) -> list[str]:
    candidates: list[str] = []
    for value in values:
        raw = str(value or "").strip()
        if not raw or raw.casefold() in _HEADERS:
            continue
        candidates.append(raw)
        if len(candidates) > MAX_CANDIDATES:
            raise ImportFormatError(f"单次最多导入 {MAX_CANDIDATES} 个候选")
    return candidates


def parse_text(text: str) -> list[str]:
    """Parse pasted comma/newline/tab/Chinese-punctuation separated values."""
    return _clean_candidates(re.split(r"[,，;；\t\r\n]+", text))


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFormatError("CSV 仅支持 UTF-8 或 GB18030 编码")


def _parse_csv(data: bytes) -> list[str]:
    return _parse_table(csv.reader(io.StringIO(_decode_csv(data))))


def _parse_table(rows: Iterable[Iterable[object]]) -> list[str]:
    table = [list(row) for row in rows]
    if not table:
        return []
    headers = [str(value or "").strip().casefold() for value in table[0]]
    code_columns = [index for index, header in enumerate(headers) if header in _CODE_HEADERS]
    name_columns = [index for index, header in enumerate(headers) if header in _NAME_HEADERS]
    if not code_columns and not name_columns:
        return _clean_candidates(cell for row in table for cell in row)

    candidates: list[object] = []
    for row in table[1:]:
        value = next(
            (
                row[index]
                for index in (*code_columns, *name_columns)
                if index < len(row) and str(row[index] or "").strip()
            ),
            None,
        )
        if value is not None:
            candidates.append(value)
    return _clean_candidates(candidates)


def _parse_xlsx(data: bytes) -> list[str]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFormatError("无法读取 XLSX 文件") from exc
    try:
        candidates = []
        for sheet in workbook.worksheets:
            candidates.extend(_parse_table([list(row) for row in sheet.iter_rows(values_only=True)]))
        return _clean_candidates(candidates)
    finally:
        workbook.close()


def parse_file(filename: str, data: bytes) -> list[str]:
    """Parse a bounded CSV/TXT/XLSX upload into search candidates."""
    if len(data) > MAX_FILE_BYTES:
        raise ImportFormatError("文件不能超过 5MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _parse_csv(data)
    if suffix == ".txt":
        return parse_text(_decode_csv(data))
    if suffix == ".xlsx":
        return _parse_xlsx(data)
    raise ImportFormatError("仅支持 CSV、TXT、XLSX 文件；图片 OCR 暂列 P3")


def _normalize_if_code(value: str) -> str | None:
    if not _DIRECT_CODE.fullmatch(value.strip()):
        return None
    try:
        return normalize_code(value)
    except ValueError:
        return None


class WatchlistImportService:
    def __init__(self, market: SearchMarket):
        self.market = market

    async def _resolve(self, candidate: str) -> dict[str, str] | None:
        direct = _normalize_if_code(candidate)
        try:
            rows = await self.market.search(candidate)
        except Exception:
            if direct:
                return {"code": direct, "name": candidate.strip().upper()}
            raise
        if rows:
            selected = next((row for row in rows if direct and row.get("code") == direct), rows[0])
            code = normalize_code(str(selected["code"]))
            return {"code": code, "name": str(selected.get("name") or code)}
        if direct:
            return {"code": direct, "name": candidate.strip().upper()}
        return None

    async def import_candidates(
        self,
        candidates: list[str],
        existing_codes: list[str] | None = None,
    ) -> dict[str, Any]:
        existing = {
            normalized
            for value in existing_codes or []
            if (normalized := _normalize_if_code(value)) is not None
        }
        imported: list[dict[str, str]] = []
        failures: list[dict[str, str]] = []
        duplicates = 0
        seen = set(existing)

        for candidate in candidates:
            try:
                item = await self._resolve(candidate)
            except Exception as exc:
                failures.append({"input": candidate, "reason": f"查询失败：{type(exc).__name__}"})
                continue
            if item is None:
                failures.append({"input": candidate, "reason": "未识别到支持的 A/HK/US 标的"})
                continue
            if item["code"] in seen:
                duplicates += 1
                continue
            seen.add(item["code"])
            imported.append(item)

        return {
            "total": len(candidates),
            "success": len(imported),
            "failed": len(failures),
            "duplicate": duplicates,
            "items": imported,
            "failures": failures,
        }
